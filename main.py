#!/usr/bin/env python3
"""
main.py — Calificador OMR (Kivy).

PASO 1: Configuración inicial (una vez) + menú con desplegables.
PASO 2: Estudiantes — lista del curso con código automático.
PASO 3: Exámenes — crear la clave de respuestas y (opcional) el PDF de la hoja.

Reutiliza gestor_salones.py y generador_pdf.py. La clave se guarda en el mismo
formato que lee lector_omr.py / registro_notas.py.
"""

from kivy.app import App
from kivy.uix.screenmanager import ScreenManager, Screen, SlideTransition
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.label import Label
from kivy.uix.button import Button
from kivy.uix.textinput import TextInput
from kivy.uix.spinner import Spinner
from kivy.uix.scrollview import ScrollView
from kivy.uix.popup import Popup
from kivy.uix.checkbox import CheckBox
from kivy.uix.image import Image
from kivy.core.window import Window
from kivy.core.text import LabelBase
from kivy.graphics import Color, RoundedRectangle, Line
from kivy.factory import Factory
from kivy.metrics import dp
from kivy.clock import Clock
from kivy.lang import Builder
from kivy.uix.anchorlayout import AnchorLayout
from kivy.uix.widget import Widget
from kivy.utils import platform
from kivy.uix.modalview import ModalView
from kivy.graphics.texture import Texture
from kivy.properties import NumericProperty, ListProperty
from kivy.animation import Animation

import os
import json
import shutil
import tempfile
import threading
from pathlib import Path
import gestor_salones as GS
import plataforma as PF
import layout_hoja as LH

BASE = GS.BASE_DIR  # carpeta "profesores"

ARCHIVO_LISTA = "lista_estudiantes.xlsx"


# ─────────────────────────────────────────────────────────────────────────────
# TEMA VISUAL · Estilo A "Institucional clásico" (solo apariencia, no lógica)
# ─────────────────────────────────────────────────────────────────────────────
AZUL     = (0.122, 0.220, 0.392, 1)   # #1F3864  (azul del escudo)
DORADO   = (0.949, 0.718, 0.020, 1)   # #F2B705  (acento dorado)
FONDO    = (0.937, 0.945, 0.961, 1)   # #EFF1F5  (fondo gris claro)
BLANCO   = (1, 1, 1, 1)
# Colores "con significado" para texto con markup (formato hexadecimal):
VERDE = "2E7D32"   # aprobó
ROJO  = "C62828"   # reprobó / error
AMBAR = "F9A825"   # por revisar

_DIR = Path(__file__).resolve().parent


def _registrar_fuentes():
    """Si hay fuentes en assets/fonts, reemplaza la fuente por defecto.
    Detecta el archivo sin importar el nombre exacto (Inter-Regular.ttf,
    Inter_18pt-Regular.ttf, etc.). Si no hay, la app sigue con la letra estándar."""
    base = _DIR / "assets" / "fonts"
    if not base.exists():
        return
    ttfs = list(base.glob("*.ttf"))
    if not ttfs:
        return

    def _buscar(*claves):
        for t in ttfs:
            n = t.name.lower()
            if all(k in n for k in claves):
                return t
        return None

    regular = _buscar("regular") or ttfs[0]
    bold = _buscar("semibold") or _buscar("bold") or regular
    try:
        LabelBase.register(name="Roboto",
                           fn_regular=str(regular),
                           fn_bold=str(bold))
    except Exception:
        pass


def _vibrar(duracion: float = 0.02):
    """Vibración corta SOLO en Android. En PC no hace nada (no rompe la prueba)."""
    try:
        if platform == "android":
            from plyer import vibrator  # type: ignore  # solo existe en Android
            vibrator.vibrate(duracion)
    except Exception:
        pass


def _solicitar_permisos_android():
    """Solo Android 9 o anterior necesita permiso para guardar en Descargas.
    Desde Android 10 no se pide ningún permiso: las fotos y planillas se eligen
    con el selector del sistema y Descargas se usa a través de MediaStore."""
    try:
        if platform != "android":
            return
        from jnius import autoclass  # type: ignore
        if autoclass("android.os.Build$VERSION").SDK_INT <= 28:
            from android.permissions import request_permissions, Permission  # type: ignore
            request_permissions([Permission.WRITE_EXTERNAL_STORAGE])
    except Exception:
        pass


class Encabezado(BoxLayout):
    """Barra azul con el escudo, el título y el subtítulo de cada pantalla."""

    def __init__(self, titulo, subtitulo="Liceo Sahagún", **kw):
        super().__init__(orientation="horizontal", size_hint_y=None,
                         height=dp(64), padding=[dp(12), dp(10)],
                         spacing=dp(12), **kw)
        with self.canvas.before:
            Color(rgba=AZUL)
            self._bg = RoundedRectangle(pos=self.pos, size=self.size,
                                        radius=[dp(12)])
        self.bind(pos=self._sync, size=self._sync)

        # Escudo "LS" (cuadro azul oscuro con borde dorado)
        badge = AnchorLayout(size_hint=(None, 1), width=dp(42))
        with badge.canvas.before:
            Color(rgba=(0.086, 0.165, 0.290, 1))
            badge._r = RoundedRectangle(pos=badge.pos, size=badge.size,
                                        radius=[dp(6), dp(6), dp(13), dp(13)])
            Color(rgba=DORADO)
            badge._l = Line(width=1.4, rounded_rectangle=(0, 0, dp(42), dp(42), dp(6)))
        badge.bind(pos=self._badge_sync, size=self._badge_sync)
        self._badge = badge
        badge.add_widget(Label(text="LS", bold=True, color=DORADO, font_size="15sp"))
        self.add_widget(badge)

        # Título + línea dorada + subtítulo
        caja = BoxLayout(orientation="vertical", spacing=dp(2))
        lbl_t = Label(text=titulo, bold=True, color=BLANCO, font_size="18sp",
                      halign="left", valign="bottom", shorten=True)
        sub = Label(text=subtitulo, color=(0.62, 0.71, 0.86, 1), font_size="12sp",
                    halign="left", valign="top")
        for lb in (lbl_t, sub):
            lb.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, i.height)))
        linea = Widget(size_hint=(None, None), size=(dp(46), dp(3)))
        with linea.canvas:
            Color(rgba=DORADO)
            linea._r = RoundedRectangle(pos=linea.pos, size=linea.size, radius=[dp(2)])
        linea.bind(pos=lambda w, *_: setattr(w._r, "pos", w.pos),
                   size=lambda w, *_: setattr(w._r, "size", w.size))
        caja.add_widget(lbl_t)
        caja.add_widget(linea)
        caja.add_widget(sub)
        self.add_widget(caja)

    def _sync(self, *_):
        self._bg.pos = self.pos
        self._bg.size = self.size

    def _badge_sync(self, *_):
        b = self._badge
        b._r.pos = b.pos
        b._r.size = b.size
        b._l.rounded_rectangle = (b.x, b.y, b.width, b.height, dp(6))


class BotonMenu(Button):
    """Botón del menú principal con un ícono vectorial dorado a la izquierda."""

    def __init__(self, icono="", **kw):
        super().__init__(**kw)
        self._icono = icono
        self.halign = "left"
        self.valign = "middle"
        self.bind(pos=self._redibujar, size=self._redibujar)

    def on_press(self):
        _vibrar()

    def _redibujar(self, *_):
        self.text_size = (self.width - dp(58), self.height)
        self.padding = [dp(50), 0, dp(10), 0]
        self.canvas.after.clear()
        s = dp(20)
        x = self.x + dp(16)
        y = self.center_y - s / 2
        with self.canvas.after:
            Color(rgba=DORADO)
            self._dibujar_icono(self._icono, x, y, s)

    @staticmethod
    def _dibujar_icono(nombre, x, y, s):
        w = 1.6
        if nombre == "estudiantes":
            Line(circle=(x + s * 0.32, y + s * 0.74, s * 0.16), width=w)
            Line(circle=(x + s * 0.70, y + s * 0.74, s * 0.16), width=w)
            Line(rounded_rectangle=(x + s * 0.08, y + s * 0.04,
                                    s * 0.46, s * 0.40, s * 0.12), width=w)
            Line(rounded_rectangle=(x + s * 0.46, y + s * 0.04,
                                    s * 0.46, s * 0.40, s * 0.12), width=w)
        elif nombre == "examenes":
            Line(rounded_rectangle=(x + s * 0.16, y, s * 0.68, s, s * 0.08), width=w)
            Line(points=[x + s * 0.30, y + s * 0.72, x + s * 0.70, y + s * 0.72], width=w)
            Line(points=[x + s * 0.30, y + s * 0.52, x + s * 0.70, y + s * 0.52], width=w)
            Line(points=[x + s * 0.30, y + s * 0.32, x + s * 0.56, y + s * 0.32], width=w)
        elif nombre == "calificar":
            Line(points=[x + s * 0.16, y + s * 0.50, x + s * 0.42, y + s * 0.22,
                         x + s * 0.86, y + s * 0.80], width=2.0)
        elif nombre == "registro":
            Line(rounded_rectangle=(x, y, s, s, s * 0.08), width=w)
            Line(points=[x, y + s * 0.5, x + s, y + s * 0.5], width=w)
            Line(points=[x + s * 0.5, y, x + s * 0.5, y + s], width=w)


class Pildora(Label):
    """Etiqueta con fondo azul claro redondeado (para 'Curso actual')."""

    def __init__(self, **kw):
        super().__init__(**kw)
        self.color = AZUL
        self.font_size = "13sp"
        self.halign = "left"
        self.valign = "middle"
        self.padding = [dp(12), 0, dp(12), 0]
        with self.canvas.before:
            Color(rgba=(0.906, 0.937, 0.984, 1))
            self._r = RoundedRectangle(pos=self.pos, size=self.size, radius=[dp(8)])
        self.bind(pos=self._upd, size=self._upd)

    def _upd(self, *_):
        self._r.pos = self.pos
        self._r.size = self.size
        self.text_size = (self.width - dp(24), self.height)


class BotonSalir(Button):
    """Botón minimalista 'Cerrar' con ícono de puerta y flecha (dibujado en canvas)."""

    def __init__(self, **kw):
        super().__init__(**kw)
        self.halign = "center"
        self.valign = "middle"
        self.bind(pos=self._redibujar, size=self._redibujar)

    def on_press(self):
        _vibrar()

    def _redibujar(self, *_):
        self.text_size = (self.width, self.height)
        self.padding = [dp(34), 0, dp(8), 0]
        self.canvas.after.clear()
        s = dp(18)
        x = self.x + dp(14)
        y = self.center_y - s / 2
        with self.canvas.after:
            Color(rgba=(0.42, 0.47, 0.55, 1))
            # Marco de la puerta
            Line(rounded_rectangle=(x, y, s * 0.55, s, s * 0.06), width=1.6)
            # Flecha saliendo hacia afuera
            Line(points=[x + s * 0.28, y + s * 0.5, x + s * 1.06, y + s * 0.5], width=1.6)
            Line(points=[x + s * 0.84, y + s * 0.74, x + s * 1.08, y + s * 0.5,
                         x + s * 0.84, y + s * 0.26], width=1.6)


class CirculoProgreso(Widget):
    """
    Anillo de progreso animado.
      • Arco dorado (grueso): cuánto se ha avanzado; crece suavemente.
      • Arco azul (fino, por dentro): gira mientras se trabaja, para que se
        note actividad aunque una foto tarde en procesarse.
      • En el centro, el porcentaje.
    """
    valor = NumericProperty(0.0)          # 0..1 (se anima con animar_a)
    giro = NumericProperty(0.0)           # ángulo del arco azul
    color_avance = ListProperty(list(DORADO))

    def __init__(self, con_porcentaje=True, **kw):
        self._con_porcentaje = con_porcentaje
        super().__init__(**kw)
        with self.canvas:
            Color(rgba=(0.906, 0.922, 0.949, 1))
            self._pista = Line(width=dp(6))
            self._c_avance = Color(rgba=self.color_avance)
            self._avance = Line(width=dp(6), cap="round")
            Color(rgba=AZUL)
            self._giro = Line(width=dp(2.2), cap="round")
        self.lbl = Label(markup=True, halign="center", valign="middle", color=AZUL)
        self.add_widget(self.lbl)
        self.bind(pos=self._dibujar, size=self._dibujar, valor=self._dibujar,
                  giro=self._dibujar, color_avance=self._dibujar)
        self._reloj = Clock.schedule_interval(self._girar, 1 / 60.0)

    def _girar(self, dt):
        self.giro = (self.giro + 260 * dt) % 360

    def _dibujar(self, *_):
        cx, cy = self.center
        r = min(self.width, self.height) / 2 - dp(8)
        self._pista.circle = (cx, cy, r)
        self._c_avance.rgba = self.color_avance
        if self.valor > 0.001:
            self._avance.circle = (cx, cy, r, 0, 360 * min(self.valor, 1.0))
        else:
            self._avance.points = []
        if self._reloj is not None:
            self._giro.circle = (cx, cy, r - dp(12), self.giro, self.giro + 80)
        else:
            self._giro.points = []
        self.lbl.size = self.size
        self.lbl.pos = self.pos
        self.lbl.text_size = self.size
        pct = int(round(self.valor * 100))
        self.lbl.text = (f"[size=34sp][b]{pct}[/b][/size][size=16sp][b] %[/b][/size]"
                         if self._con_porcentaje else "")

    def animar_a(self, valor):
        Animation.cancel_all(self, "valor")
        Animation(valor=valor, d=0.45, t="out_cubic").start(self)

    def detener(self):
        """Deja de girar el arco azul (el trabajo terminó)."""
        if self._reloj is not None:
            self._reloj.cancel()
            self._reloj = None
        self._dibujar()


class _VentanaSinRastro(ModalView):
    """ModalView que fuerza el repintado al cerrarse: si no, en algunos
    equipos la tarjeta quedaba "pegada" en pantalla hasta el siguiente toque."""

    def __init__(self, **kw):
        super().__init__(**kw)
        self.bind(_is_open=self._al_cerrarse)

    @staticmethod
    def _al_cerrarse(_inst, abierto):
        if not abierto:
            Clock.schedule_once(lambda dt: Window.canvas.ask_update(), 0)


def _tarjeta_blanca(**kw):
    tarjeta = BoxLayout(orientation="vertical", **kw)
    with tarjeta.canvas.before:
        Color(rgba=BLANCO)
        fondo = RoundedRectangle(radius=[dp(18)])
    tarjeta.bind(pos=lambda w, *_: setattr(fondo, "pos", w.pos),
                 size=lambda w, *_: setattr(fondo, "size", w.size))
    return tarjeta


class PopupEspera(_VentanaSinRastro):
    """Ventana "Un momento…" con el anillo girando (sin porcentaje)."""

    def __init__(self, mensaje, **kw):
        super().__init__(size_hint=(0.8, None), height=dp(250), auto_dismiss=False,
                         background_color=(0, 0, 0, 0),
                         overlay_color=(0.05, 0.08, 0.15, 0.6), **kw)
        tarjeta = _tarjeta_blanca(padding=[dp(20), dp(18)], spacing=dp(10))
        caja = AnchorLayout(size_hint_y=None, height=dp(130))
        circulo = CirculoProgreso(size_hint=(None, None), size=(dp(120), dp(120)),
                                  con_porcentaje=False)
        caja.add_widget(circulo)
        lbl = Label(text=mensaje, bold=True, color=AZUL, font_size="16sp",
                    halign="center", valign="middle")
        lbl.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, None)))
        tarjeta.add_widget(caja)
        tarjeta.add_widget(lbl)
        self.add_widget(tarjeta)


class PopupProgreso(_VentanaSinRastro):
    """Ventana de "Calificando…" con anillo de progreso, contadores y Cancelar."""

    def __init__(self, titulo, subtitulo, total, al_cancelar, **kw):
        super().__init__(size_hint=(0.9, None), height=dp(440), auto_dismiss=False,
                         background_color=(0, 0, 0, 0),
                         overlay_color=(0.05, 0.08, 0.15, 0.6), **kw)
        self.total = total
        self._al_cancelar = al_cancelar
        self._cuenta = {"ok": 0, "revisar": 0, "error": 0}

        tarjeta = _tarjeta_blanca(padding=[dp(22), dp(20)], spacing=dp(6))

        def etiqueta(texto, alto, **kw):
            lb = Label(text=texto, size_hint_y=None, height=dp(alto), halign="center",
                       valign="middle", shorten=True, shorten_from="center", **kw)
            lb.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, i.height)))
            return lb

        self.lbl_titulo = etiqueta(titulo, 28, bold=True, color=AZUL, font_size="19sp")
        self.lbl_sub = etiqueta(subtitulo, 20, color=(0.42, 0.47, 0.55, 1), font_size="13sp")
        caja = AnchorLayout(size_hint_y=None, height=dp(190))
        self.circulo = CirculoProgreso(size_hint=(None, None), size=(dp(176), dp(176)))
        caja.add_widget(self.circulo)
        self.lbl_estado = etiqueta("Preparando…", 24, bold=True, font_size="16sp")
        self.lbl_archivo = etiqueta("", 18, color=(0.42, 0.47, 0.55, 1), font_size="12sp")
        self.lbl_cuenta = etiqueta("", 22, markup=True, font_size="13sp")
        self.zona_boton = BoxLayout(size_hint_y=None, height=dp(48))
        self.btn = Factory.BotonOutline(text="Cancelar")
        self.btn.bind(on_release=self._cancelar)
        self.zona_boton.add_widget(self.btn)

        for w in (self.lbl_titulo, self.lbl_sub, caja, self.lbl_estado,
                  self.lbl_archivo, self.lbl_cuenta, Widget(size_hint_y=None, height=dp(4)),
                  self.zona_boton):
            tarjeta.add_widget(w)
        self.add_widget(tarjeta)
        self._mostrar_cuenta()

    def _mostrar_cuenta(self):
        c = self._cuenta
        self.lbl_cuenta.text = (f"[color={VERDE}]{c['ok']} listas[/color]   ·   "
                                f"[color={AMBAR}]{c['revisar']} por revisar[/color]   ·   "
                                f"[color={ROJO}]{c['error']} con error[/color]")

    def _cancelar(self, *_):
        self.btn.disabled = True
        self.btn.text = "Cancelando…"
        self.lbl_estado.text = "Terminando la hoja actual…"
        self._al_cancelar()

    def empezar(self, i, nombre):
        """Se empieza a procesar la hoja i (0-based)."""
        if not self.btn.disabled:
            self.lbl_estado.text = f"Calificando hoja {i + 1} de {self.total}"
        self.lbl_archivo.text = nombre

    def avance(self, hechas, estado):
        """Terminó una hoja. estado: 'ok', 'revisar' o 'error'."""
        self._cuenta[estado] += 1
        self._mostrar_cuenta()
        self.circulo.animar_a(hechas / float(self.total))

    def terminar(self, cuenta, cancelado, al_cerrar):
        """Muestra el resumen final y cambia el botón a 'Ver resultados'."""
        self._cuenta = cuenta
        self._mostrar_cuenta()
        self.circulo.detener()
        hechas = sum(cuenta.values())
        if cancelado:
            self.lbl_titulo.text = "Calificación cancelada"
            self.lbl_estado.text = f"Se calificaron {hechas} de {self.total} hojas"
        else:
            self.lbl_titulo.text = "¡Listo!"
            self.circulo.animar_a(1.0)
            self.lbl_estado.text = (f"{hechas} hoja calificada" if hechas == 1
                                    else f"{hechas} hojas calificadas")
        self.lbl_archivo.text = ("Revisa las hojas marcadas en amarillo o rojo."
                                 if cuenta["revisar"] or cuenta["error"] else "")
        # Anillo verde si todo salió bien; ámbar si hay algo que revisar
        self.circulo.color_avance = list(
            (0.180, 0.490, 0.196, 1) if not (cuenta["revisar"] or cuenta["error"])
            else (0.976, 0.659, 0.145, 1))
        self.zona_boton.clear_widgets()
        btn = Factory.BotonAccion(text="Ver resultados")
        btn.bind(on_release=lambda *_: (self.dismiss(), al_cerrar()))
        self.zona_boton.add_widget(btn)


def _caption(texto):
    """Etiqueta pequeña gris que va encima de un campo."""
    lb = Label(text=texto, color=(0.42, 0.47, 0.55, 1), font_size="12sp",
               size_hint_y=None, height=dp(18), halign="left", valign="middle")
    lb.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, i.height)))
    return lb


# Tema visual (embebido para no depender de un archivo .kv externo) -------------
_TEMA_KV = """
#:set C_AZUL      (0.122, 0.220, 0.392, 1)
#:set C_AZUL_OSC  (0.078, 0.157, 0.298, 1)
#:set C_DORADO    (0.949, 0.718, 0.020, 1)
#:set C_DORADO_OSC (0.984, 0.788, 0.290, 1)
#:set C_DORADO_TXT (0.227, 0.169, 0.0, 1)
#:set C_ROJO      (0.780, 0.160, 0.160, 1)
#:set C_ROJO_OSC  (0.620, 0.120, 0.120, 1)
#:set C_TEXTO     (0.106, 0.129, 0.176, 1)
#:set C_GRIS      (0.42, 0.47, 0.55, 1)
#:set C_BORDE     (0.79, 0.83, 0.89, 1)
#:set C_BLANCO    (1, 1, 1, 1)

<Label>:
    color: C_TEXTO
    font_size: '15sp'

<Button>:
    background_normal: ''
    background_down: ''
    background_disabled_normal: ''
    background_color: 0, 0, 0, 0
    color: C_BLANCO
    font_size: '15sp'
    canvas.before:
        Color:
            rgba: C_AZUL if self.state == 'normal' else C_AZUL_OSC
        RoundedRectangle:
            pos: self.pos
            size: self.size
            radius: [10,]

<BotonAccion@Button>:
    color: C_DORADO_TXT
    bold: True
    canvas.before:
        Color:
            rgba: C_DORADO if self.state == 'normal' else C_DORADO_OSC
        RoundedRectangle:
            pos: self.pos
            size: self.size
            radius: [10,]

<BotonPeligro@Button>:
    color: C_BLANCO
    canvas.before:
        Color:
            rgba: C_ROJO if self.state == 'normal' else C_ROJO_OSC
        RoundedRectangle:
            pos: self.pos
            size: self.size
            radius: [10,]

<BotonOutline@Button>:
    color: C_DORADO_TXT
    bold: True
    canvas.before:
        Color:
            rgba: (1, 1, 1, 1) if self.state == 'normal' else (0.98, 0.96, 0.88, 1)
        RoundedRectangle:
            pos: self.pos
            size: self.size
            radius: [10,]
        Color:
            rgba: C_DORADO
        Line:
            rounded_rectangle: (self.x + 1, self.y + 1, self.width - 2, self.height - 2, 10)
            width: 1.4

<FilaLista@Button>:
    color: C_TEXTO
    markup: True
    halign: 'left'
    valign: 'middle'
    padding: [dp(12), 0]
    text_size: self.size
    font_size: '15sp'
    canvas.before:
        Color:
            rgba: (1, 1, 1, 1) if self.state == 'normal' else (0.94, 0.95, 0.97, 1)
        RoundedRectangle:
            pos: self.pos
            size: self.size
            radius: [8,]

<BotonSalir>:
    color: C_GRIS
    bold: False
    canvas.before:
        Color:
            rgba: (1, 1, 1, 1) if self.state == 'normal' else (0.95, 0.96, 0.98, 1)
        RoundedRectangle:
            pos: self.pos
            size: self.size
            radius: [10,]
        Color:
            rgba: C_BORDE
        Line:
            rounded_rectangle: (self.x + 1, self.y + 1, self.width - 2, self.height - 2, 10)
            width: 1.2

<Spinner>:
    background_normal: ''
    background_down: ''
    background_color: 0, 0, 0, 0
    color: C_AZUL
    font_size: '14sp'
    canvas.before:
        Color:
            rgba: C_BLANCO
        RoundedRectangle:
            pos: self.pos
            size: self.size
            radius: [10,]
    canvas.after:
        Color:
            rgba: C_BORDE
        Line:
            rounded_rectangle: (self.x + 1, self.y + 1, self.width - 2, self.height - 2, 10)
            width: 1.2

<TextInput>:
    background_color: C_BLANCO
    foreground_color: C_TEXTO
    cursor_color: C_DORADO
    hint_text_color: C_GRIS
    font_size: '15sp'
    padding: [12, 12, 12, 12]
    canvas.after:
        Color:
            rgba: C_DORADO if self.focus else C_BORDE
        Line:
            rounded_rectangle: (self.x + 1, self.y + 1, self.width - 2, self.height - 2, 8)
            width: 1.3

<Popup>:
    background_color: 0, 0, 0, 0
    title_color: C_AZUL
    title_size: '17sp'
    separator_color: C_DORADO
    separator_height: '2dp'
    canvas.before:
        Color:
            rgba: C_BLANCO
        RoundedRectangle:
            pos: self.pos
            size: self.size
            radius: [14,]
"""

Builder.load_string(_TEMA_KV)


def _subdirs(p):
    return sorted(d.name for d in p.iterdir() if d.is_dir()) if p.exists() else []


def aviso(titulo, mensaje):
    _popup_texto(titulo, mensaje, [("Cerrar", "normal", None)])


def confirmar(titulo, mensaje, texto_si, al_si, peligro=True):
    """Pregunta antes de algo que no se puede deshacer."""
    _popup_texto(titulo, mensaje, [("Cancelar", "outline", None),
                                   (texto_si, "peligro" if peligro else "accion", al_si)])


def aviso_archivo(titulo, mensaje, ref, nombre):
    """Aviso de archivo guardado, con botones para abrirlo o compartirlo."""
    def _abrir():
        err = PF.abrir(ref, nombre)
        if err:
            aviso("No se pudo abrir", f"Ábrelo desde Descargas.\n\n({err})")

    def _compartir():
        err = PF.compartir(ref, nombre)
        if err:
            aviso("No se pudo compartir", err)

    _popup_texto(titulo, mensaje, [("Abrir", "accion", _abrir),
                                   ("Compartir" if PF.ES_ANDROID else "Ver carpeta", "normal", _compartir),
                                   ("Cerrar", "outline", None)])


def _popup_texto(titulo, mensaje, botones, alto=0.55):
    """Ventana con un texto y una fila de botones [(texto, estilo, acción)].
    estilo: 'accion' (dorado), 'peligro' (rojo), 'outline' o 'normal' (azul)."""
    cont = BoxLayout(orientation="vertical", padding=dp(10), spacing=dp(10))
    lbl = Label(text=str(mensaje), halign="center", valign="top", size_hint_y=None)
    lbl.bind(width=lambda i, w: setattr(i, "text_size", (w, None)))
    lbl.bind(texture_size=lambda i, ts: setattr(i, "height", ts[1]))
    scroll = ScrollView()
    scroll.add_widget(lbl)
    cont.add_widget(scroll)
    fila = BoxLayout(size_hint_y=None, height=dp(46), spacing=dp(8))
    popup = Popup(title=titulo, content=cont, size_hint=(0.92, alto))
    clases = {"accion": Factory.BotonAccion, "peligro": Factory.BotonPeligro,
              "outline": Factory.BotonOutline, "normal": Button}
    for texto, estilo, accion in botones:
        btn = clases[estilo](text=texto, font_size="14sp")

        def _pulsar(*_, a=accion):
            popup.dismiss()
            if a is not None:
                a()
        btn.bind(on_release=_pulsar)
        fila.add_widget(btn)
    cont.add_widget(fila)
    popup.open()
    return popup


def en_segundo_plano(mensaje, trabajo, al_terminar):
    """
    Corre trabajo() en otro hilo mostrando una ventana de espera (la interfaz
    no se congela). Al final llama a al_terminar(resultado, error) en el hilo
    de la interfaz; error es None si todo salió bien.
    """
    espera = PopupEspera(mensaje)
    espera.open()

    def _hilo():
        try:
            res, err = trabajo(), None
        except Exception as e:
            res, err = None, e

        def _fin(dt):
            espera.dismiss()
            al_terminar(res, err)
        Clock.schedule_once(_fin, 0)
    threading.Thread(target=_hilo, daemon=True).start()


def codigo_de(posicion: int) -> str:
    return f"{posicion:03d}"


def nota_local(aciertos, total):
    """Nota 0-10 truncada a 1 decimal, con piso 1.0 (igual que el lector)."""
    if not total:
        return 0.0
    nota = 10.0 * aciertos / total
    nota = int(nota * 10 + 1e-9) / 10.0
    return max(1.0, nota)


# ─────────────────────────────────────────────────────────────────────────────
# Planilla de estudiantes (mismo formato que registro_notas.py)
# ─────────────────────────────────────────────────────────────────────────────
def leer_lista_estudiantes(ruta_curso):
    archivo = ruta_curso / "estudiantes" / ARCHIVO_LISTA
    if not archivo.exists():
        return []
    from openpyxl import load_workbook
    ws = load_workbook(archivo).active
    nombres = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if len(fila) > 2 and fila[2] and str(fila[2]).strip():
            nombres.append(str(fila[2]).strip())
    return nombres


def guardar_lista_estudiantes(ruta_curso, nombres):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    carpeta = ruta_curso / "estudiantes"
    carpeta.mkdir(parents=True, exist_ok=True)
    archivo = carpeta / ARCHIVO_LISTA
    wb = Workbook()
    ws = wb.active
    ws.title = "Estudiantes"
    ws.append(["N° Lista", "Código", "Nombre"])
    for c in range(1, 4):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1F3864")
        cell.alignment = Alignment(horizontal="center")
    for pos, nombre in enumerate(nombres, start=1):
        ws.append([pos, codigo_de(pos), nombre])
        cod = ws.cell(row=pos + 1, column=2)
        cod.number_format = "@"
        cod.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 34
    wb.save(archivo)
    return archivo


def _openpyxl_disponible():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def importar_planilla(ruta_archivo):
    """
    Lee los nombres de los estudiantes desde un archivo .xlsx o .csv.
    - Si hay una columna con encabezado que contenga 'nombre', usa esa.
    - Si no, usa la columna con más texto (nombres).
    Devuelve una lista de nombres (en orden de aparición).
    """
    ruta = Path(ruta_archivo)
    ext = ruta.suffix.lower()
    filas = []
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(ruta, data_only=True).active
        filas = [list(r) for r in ws.iter_rows(values_only=True)]
    elif ext == ".csv":
        import csv as _csv
        with open(ruta, newline="", encoding="utf-8-sig") as f:
            filas = [row for row in _csv.reader(f)]
    else:
        return []

    filas = [f for f in filas if f and any(c not in (None, "") for c in f)]
    if not filas:
        return []

    # Buscar columna por encabezado "nombre"
    encabezado = [str(c).strip().lower() if c is not None else "" for c in filas[0]]
    col = next((i for i, h in enumerate(encabezado) if "nombre" in h), None)
    inicio = 1 if col is not None else 0

    if col is None:
        # Sin encabezado claro: columna con más texto no numérico
        ncols = max(len(f) for f in filas)
        mejor, mejor_count = 0, -1
        for i in range(ncols):
            count = sum(1 for f in filas
                        if i < len(f) and isinstance(f[i], str)
                        and f[i].strip() and not f[i].strip().isdigit())
            if count > mejor_count:
                mejor, mejor_count = i, count
        col = mejor

    nombres = []
    for f in filas[inicio:]:
        if col < len(f) and f[col] is not None:
            v = str(f[col]).strip()
            if v and v.lower() != "nombre":
                nombres.append(v)
    return nombres


# ─────────────────────────────────────────────────────────────────────────────
# Exámenes (clave JSON + PDF opcional)
# ─────────────────────────────────────────────────────────────────────────────
def listar_examenes(ruta_curso):
    carpeta = ruta_curso / "examenes_claves"
    if not carpeta.exists():
        return []
    out = []
    for f in sorted(carpeta.glob("*.json")):
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
            # Las claves viejas no traen formato: todas eran de hoja completa
            formato = LH.obtener(d.get("formato")).clave
            out.append((f, d.get("examen", f.stem), d.get("total_preguntas", "?"), formato))
        except Exception:
            continue
    return out


def guardar_clave(ruta_curso, nombre, cantidad, clave_letras, formato):
    carpeta = ruta_curso / "examenes_claves"
    carpeta.mkdir(parents=True, exist_ok=True)
    nombre = nombre.replace(" ", "_")
    datos = {"examen": nombre, "total_preguntas": cantidad,
             "clave_correctas": clave_letras, "formato": LH.obtener(formato).clave}
    archivo = carpeta / f"respuestas_{nombre}.json"
    archivo.write_text(json.dumps(datos, indent=4, ensure_ascii=False), encoding="utf-8")
    return archivo


def generar_pdf_a_descargas(nombre, cantidad, formato="carta"):
    """Genera la hoja PDF y la deja en Descargas. Devuelve (texto, referencia,
    nombre_archivo); la referencia sirve para abrirla o compartirla."""
    import generador_pdf as GP  # importación tardía
    nombre_arch = f"{nombre.replace(' ', '_')}_hoja.pdf"
    tmp = Path(tempfile.mkdtemp(prefix="hoja_")) / nombre_arch
    GP.generar_hoja_omr(tmp, cantidad, formato)
    texto, ref = PF.guardar_en_descargas(tmp, nombre_arch)
    shutil.rmtree(tmp.parent, ignore_errors=True)
    return texto, ref, nombre_arch


# ─────────────────────────────────────────────────────────────────────────────
# PANTALLA: CONFIGURACIÓN INICIAL
# ─────────────────────────────────────────────────────────────────────────────
class SetupScreen(Screen):
    def __init__(self, **kw):
        super().__init__(**kw)
        root = BoxLayout(orientation="vertical", padding=dp(18), spacing=dp(12))
        root.add_widget(Encabezado("Configuración inicial"))

        intro = Label(text="Esto se hace una sola vez.", size_hint_y=None,
                      height=dp(24), font_size="13sp", color=(0.42, 0.47, 0.55, 1),
                      halign="center", valign="middle")
        intro.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, i.height)))
        root.add_widget(intro)

        root.add_widget(_caption("¿Cómo te llamas?"))
        self.in_nombre = TextInput(hint_text="Nombre y apellido", multiline=False,
                                   size_hint_y=None, height=dp(50))
        root.add_widget(self.in_nombre)

        root.add_widget(_caption("Tus grados y cuántos cursos tiene cada uno"))
        self.filas_box = BoxLayout(orientation="vertical", size_hint_y=None,
                                   spacing=dp(10))
        self.filas_box.bind(minimum_height=self.filas_box.setter("height"))
        scroll = ScrollView()
        scroll.add_widget(self.filas_box)
        root.add_widget(scroll)

        self.filas = []
        self._agregar_fila()

        btn_add = Factory.BotonOutline(text="+ Agregar otro grado", size_hint_y=None,
                                       height=dp(50))
        btn_add.bind(on_release=lambda *_: self._agregar_fila())
        root.add_widget(btn_add)

        btn_crear = Factory.BotonAccion(text="Crear y continuar", size_hint_y=None,
                                        height=dp(56))
        btn_crear.bind(on_release=self._crear)
        root.add_widget(btn_crear)
        self.add_widget(root)

    def _agregar_fila(self):
        fila = BoxLayout(size_hint_y=None, height=dp(74), spacing=dp(12))

        col_g = BoxLayout(orientation="vertical", spacing=dp(4))
        col_g.add_widget(_caption("Grado"))
        in_grado = TextInput(hint_text="ej. 11", multiline=False, input_filter="int",
                             size_hint_y=None, height=dp(48))
        col_g.add_widget(in_grado)

        col_c = BoxLayout(orientation="vertical", spacing=dp(4), size_hint_x=0.55)
        col_c.add_widget(_caption("Cursos"))
        sp = Spinner(text="1", values=[str(i) for i in range(1, 11)],
                     size_hint_y=None, height=dp(48))
        col_c.add_widget(sp)

        fila.add_widget(col_g)
        fila.add_widget(col_c)
        self.filas_box.add_widget(fila)
        self.filas.append((in_grado, sp))

    def _crear(self, *_):
        nombre = self.in_nombre.text.strip()
        if not nombre:
            aviso("Falta el nombre", "Escribe tu nombre para continuar.")
            return
        grados_cursos = []
        for in_grado, sp in self.filas:
            g = in_grado.text.strip()
            if g:
                grados_cursos.append((g, int(sp.text)))
        if not grados_cursos:
            aviso("Faltan grados", "Agrega al menos un grado con sus cursos.")
            return
        GS.inicializar_sistema()
        ruta_prof = GS.crear_profesor(nombre)
        for grado, n in grados_cursos:
            ruta_grado = GS.crear_grado(ruta_prof, grado)
            for i in range(n):
                GS.crear_curso(ruta_grado, chr(65 + i))
        App.get_running_app().ir_a_home()


# ─────────────────────────────────────────────────────────────────────────────
# PANTALLA: MENÚ PRINCIPAL
# ─────────────────────────────────────────────────────────────────────────────
class HomeScreen(Screen):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.contenedor = BoxLayout(orientation="vertical", padding=dp(14), spacing=dp(8))
        self.add_widget(self.contenedor)

    def on_pre_enter(self, *_):
        self.construir()

    def construir(self):
        self.contenedor.clear_widgets()
        b = self.contenedor
        b.add_widget(Encabezado("Calificador OMR"))

        profesores = _subdirs(BASE)
        b.add_widget(_caption("Profesor"))
        self.sp_prof = Spinner(text=profesores[0] if profesores else "",
                               values=profesores, size_hint_y=None, height=dp(46))
        self.sp_prof.bind(text=lambda *_: self._refrescar_grados())
        b.add_widget(self.sp_prof)

        fila_gc = BoxLayout(size_hint_y=None, height=dp(70), spacing=dp(10))
        col_g = BoxLayout(orientation="vertical", spacing=dp(4))
        col_g.add_widget(_caption("Grado"))
        self.sp_grado = Spinner(text="", values=[], size_hint_y=None, height=dp(46))
        self.sp_grado.bind(text=lambda *_: self._refrescar_cursos())
        col_g.add_widget(self.sp_grado)
        col_c = BoxLayout(orientation="vertical", spacing=dp(4))
        col_c.add_widget(_caption("Curso"))
        self.sp_curso = Spinner(text="", values=[], size_hint_y=None, height=dp(46))
        self.sp_curso.bind(text=lambda *_: self._guardar_seleccion())
        col_c.add_widget(self.sp_curso)
        fila_gc.add_widget(col_g)
        fila_gc.add_widget(col_c)
        b.add_widget(fila_gc)

        self.lbl_actual = Pildora(text="", size_hint_y=None, height=dp(32))
        b.add_widget(self.lbl_actual)

        b.add_widget(Widget())  # espacio flexible: empuja el botón principal abajo

        btn_cal = Factory.BotonAccion(text="Calificar", size_hint_y=None, height=dp(54))
        btn_cal.bind(on_press=lambda *_: _vibrar())
        btn_cal.bind(on_release=lambda *_: self._entrar())
        b.add_widget(btn_cal)

        btn_add = Factory.BotonOutline(text="+ Agregar grado o curso",
                                       size_hint_y=None, height=dp(46))
        btn_add.bind(on_release=lambda *_: self._popup_agregar())
        b.add_widget(btn_add)

        btn_resp = Button(text="Copia de seguridad", size_hint_y=None, height=dp(44))
        btn_resp.bind(on_release=lambda *_: self._popup_respaldo())
        b.add_widget(btn_resp)

        self._refrescar_grados()

    # ── Copia de seguridad ──────────────────────────────────────────────────
    def _popup_respaldo(self):
        _popup_texto(
            "Copia de seguridad",
            "Guarda TODO (listas, claves, notas e imágenes) en un archivo .zip en "
            "Descargas. Súbelo a Drive o envíatelo: si cambias de celular o "
            "desinstalas la app, con ese archivo recuperas todo.",
            [("Guardar copia", "accion", self._guardar_respaldo),
             ("Restaurar", "normal", self._elegir_respaldo),
             ("Cerrar", "outline", None)], alto=0.5)

    def _guardar_respaldo(self):
        def trabajo():
            zip_path = PF.crear_respaldo(BASE)
            try:
                return PF.guardar_en_descargas(zip_path, zip_path.name) + (zip_path.name,)
            finally:
                zip_path.unlink()

        def listo(res, err):
            if err:
                aviso("No se pudo", f"No se pudo crear la copia:\n{err}")
                return
            texto, ref, nombre = res
            aviso_archivo("Copia guardada",
                          f"Copia guardada en:\n{texto}\n\nSúbela a Drive o envíatela "
                          "para tenerla fuera del celular.", ref, nombre)
        en_segundo_plano("Guardando copia\u2026", trabajo, listo)

    def _elegir_respaldo(self):
        PF.elegir_archivos(self._respaldo_elegido, tipo="zip", multiple=False,
                           titulo="Elige la copia de seguridad (.zip)")

    def _respaldo_elegido(self, refs):
        nombre = PF.nombre_de(refs[0])
        if not nombre.lower().endswith(".zip"):
            aviso("Archivo no válido", f"'{nombre}' no es una copia de seguridad (.zip).")
            return

        def trabajo():
            zip_local = PF.copiar_a_temporal(refs[0])
            try:
                return PF.restaurar_respaldo(zip_local, BASE)
            finally:
                shutil.rmtree(zip_local.parent, ignore_errors=True)

        def listo(anterior, err):
            if err:
                aviso("No se pudo restaurar", str(err))
                return
            App.get_running_app().ruta_curso = None
            self.construir()
            aviso("Copia restaurada",
                  "Se restauraron los datos de la copia.\n\n"
                  f"Lo que había antes no se borró: quedó guardado aparte ({anterior.name}).")

        confirmar("Restaurar copia",
                  f"Los datos actuales se reemplazarán por los de '{nombre}'.\n"
                  "(Lo que hay ahora queda guardado aparte, por si acaso.)",
                  "Restaurar", lambda: en_segundo_plano("Restaurando copia\u2026", trabajo, listo))

    def _entrar(self):
        app = App.get_running_app()
        if app.ruta_curso is None:
            aviso("Falta el curso", "Selecciona grado y curso primero.")
            return
        app.sm.current = "dashboard"

    def _refrescar_grados(self):
        prof = self.sp_prof.text
        grados = _subdirs(BASE / prof) if prof else []
        self.sp_grado.values = grados
        self.sp_grado.text = grados[0] if grados else ""
        self._refrescar_cursos()

    def _refrescar_cursos(self):
        prof, grado = self.sp_prof.text, self.sp_grado.text
        cursos = _subdirs(BASE / prof / grado) if (prof and grado) else []
        self.sp_curso.values = cursos
        self.sp_curso.text = cursos[0] if cursos else ""
        self._guardar_seleccion()

    def _guardar_seleccion(self):
        app = App.get_running_app()
        app.profesor = self.sp_prof.text
        app.grado = self.sp_grado.text
        app.curso = self.sp_curso.text
        if app.profesor and app.grado and app.curso:
            app.ruta_curso = BASE / app.profesor / app.grado / app.curso
            self.lbl_actual.text = (f"Curso actual: {app.grado}°{app.curso}"
                                    f"  ·  {app.profesor.replace('_', ' ')}")
        else:
            app.ruta_curso = None
            self.lbl_actual.text = "Selecciona grado y curso."

    def _popup_agregar(self):
        cont = BoxLayout(orientation="vertical", spacing=dp(8), padding=dp(10))
        cont.add_widget(Label(text="Agregar un grado nuevo (con sus cursos)",
                              size_hint_y=None, height=dp(28)))
        fila = BoxLayout(size_hint_y=None, height=dp(44), spacing=dp(6))
        fila.add_widget(Label(text="Grado:", size_hint_x=0.35))
        in_g = TextInput(hint_text="ej. 10", multiline=False, input_filter="int")
        fila.add_widget(in_g)
        cont.add_widget(fila)
        fila2 = BoxLayout(size_hint_y=None, height=dp(44), spacing=dp(6))
        fila2.add_widget(Label(text="# Cursos:", size_hint_x=0.35))
        sp = Spinner(text="1", values=[str(i) for i in range(1, 11)])
        fila2.add_widget(sp)
        cont.add_widget(fila2)
        btn = Button(text="Crear", size_hint_y=None, height=dp(46))
        cont.add_widget(btn)
        pop = Popup(title="Agregar grado o curso", content=cont, size_hint=(0.9, 0.6))

        def _crear(*_):
            g = in_g.text.strip()
            if not g:
                return
            ruta_prof = BASE / self.sp_prof.text
            ruta_grado = GS.crear_grado(ruta_prof, g)
            inicio = len(_subdirs(ruta_grado))
            for i in range(int(sp.text)):
                GS.crear_curso(ruta_grado, chr(65 + inicio + i))
            pop.dismiss()
            self.construir()

        btn.bind(on_release=_crear)
        pop.open()


# ─────────────────────────────────────────────────────────────────────────────
# PANTALLA: PANEL DEL CURSO (Dashboard)
# ─────────────────────────────────────────────────────────────────────────────
class DashboardScreen(Screen):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.contenedor = BoxLayout(orientation="vertical", padding=dp(14), spacing=dp(8))
        self.add_widget(self.contenedor)

    def on_pre_enter(self, *_):
        self.construir()

    def construir(self):
        app = App.get_running_app()
        b = self.contenedor
        b.clear_widgets()
        b.add_widget(Encabezado("Panel del curso"))

        # Banner persistente "Actualmente estás calificando: Grado - Curso"
        banner = BoxLayout(size_hint_y=None, height=dp(48), padding=[dp(12), dp(6)])
        with banner.canvas.before:
            Color(rgba=DORADO)
            banner._r = RoundedRectangle(pos=banner.pos, size=banner.size,
                                         radius=[dp(10)])
        banner.bind(pos=lambda w, *_: setattr(w._r, "pos", w.pos),
                    size=lambda w, *_: setattr(w._r, "size", w.size))
        lbl = Label(markup=True, color=(0.227, 0.169, 0.0, 1),
                    halign="center", valign="middle", font_size="15sp",
                    text=f"Actualmente estás calificando:  [b]{app.grado}\u00b0 {app.curso}[/b]")
        lbl.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, i.height)))
        banner.add_widget(lbl)
        b.add_widget(banner)

        acciones = {
            "Estudiantes": ("estudiantes", "estudiantes"),
            "Exámenes": ("examenes", "examenes"),
            "Calificar": ("calificar", "calificar"),
            "Registro de notas": ("registro", "registro"),
        }
        for texto, (icono, destino) in acciones.items():
            btn = BotonMenu(icono=icono, text=texto, size_hint_y=None, height=dp(56))
            btn.bind(on_release=lambda inst, d=destino: self._ir(d))
            b.add_widget(btn)

        b.add_widget(Widget())  # espacio flexible

        btn_cambiar = Button(text="\u2190 Cambiar curso", size_hint_y=None, height=dp(46))
        btn_cambiar.bind(on_release=lambda *_: setattr(app.sm, "current", "home"))
        b.add_widget(btn_cambiar)

        btn_salir = BotonSalir(text="Cerrar", size_hint_y=None, height=dp(46))
        btn_salir.bind(on_release=lambda *_: App.get_running_app().stop())
        b.add_widget(btn_salir)

    def _ir(self, pantalla):
        app = App.get_running_app()
        if app.ruta_curso is None:
            aviso("Falta el curso", "Vuelve y selecciona grado y curso.")
            return
        app.sm.current = pantalla


# ─────────────────────────────────────────────────────────────────────────────
# PANTALLA: ESTUDIANTES
# ─────────────────────────────────────────────────────────────────────────────
class EstudiantesScreen(Screen):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.cont = BoxLayout(orientation="vertical", padding=dp(14), spacing=dp(8))
        self.add_widget(self.cont)

    def on_pre_enter(self, *_):
        self.construir()

    def construir(self):
        app = App.get_running_app()
        self.cont.clear_widgets()
        b = self.cont
        b.add_widget(Encabezado(f"Estudiantes · {app.grado}°{app.curso}"))
        if not _openpyxl_disponible():
            b.add_widget(Label(text="Falta la librería 'openpyxl'.\nInstala: pip install openpyxl"))
            self._boton_volver(b)
            return
        nombres = leer_lista_estudiantes(app.ruta_curso)
        b.add_widget(_caption(f"{len(nombres)} estudiantes  ·  toca uno para editarlo"
                              if nombres else "Lista del curso"))
        lista_box = BoxLayout(orientation="vertical", size_hint_y=None, spacing=dp(4))
        lista_box.bind(minimum_height=lista_box.setter("height"))
        if nombres:
            for pos, nombre in enumerate(nombres, start=1):
                fila = Factory.FilaLista(text=f"[b]{codigo_de(pos)}[/b]     {nombre}",
                                         size_hint_y=None, height=dp(40))
                fila.bind(on_release=lambda inst, p=pos: self._editar(p))
                lista_box.add_widget(fila)
        else:
            lista_box.add_widget(Label(text="(Aún no hay estudiantes)",
                                       size_hint_y=None, height=dp(30)))
        scroll = ScrollView()
        scroll.add_widget(lista_box)
        b.add_widget(scroll)
        fila = BoxLayout(size_hint_y=None, height=dp(46), spacing=dp(6))
        self.in_nombre = TextInput(hint_text="Nombre del estudiante", multiline=False)
        self.in_nombre.bind(on_text_validate=self._agregar)
        fila.add_widget(self.in_nombre)
        btn_add = Button(text="Agregar", size_hint_x=0.3)
        btn_add.bind(on_release=self._agregar)
        fila.add_widget(btn_add)
        b.add_widget(fila)
        fila2 = BoxLayout(size_hint_y=None, height=dp(44), spacing=dp(6))
        btn_importar = Button(text="Importar Excel/CSV")
        btn_importar.bind(on_release=self._importar)
        fila2.add_widget(btn_importar)
        btn_vaciar = Factory.BotonPeligro(text="Vaciar lista", size_hint_x=0.45)
        btn_vaciar.bind(on_release=self._vaciar)
        fila2.add_widget(btn_vaciar)
        b.add_widget(fila2)
        self._boton_volver(b)

    def _boton_volver(self, b):
        btn = Button(text="← Volver al menú", size_hint_y=None, height=dp(44))
        btn.bind(on_release=lambda *_: setattr(App.get_running_app().sm, "current", "dashboard"))
        b.add_widget(btn)

    def _agregar(self, *_):
        app = App.get_running_app()
        nombre = self.in_nombre.text.strip()
        if not nombre:
            return
        nombres = leer_lista_estudiantes(app.ruta_curso)
        nombres.append(nombre)
        guardar_lista_estudiantes(app.ruta_curso, nombres)
        self.construir()

    def _editar(self, pos):
        """Editar el nombre o eliminar al estudiante de la posición 'pos' (1-based)."""
        app = App.get_running_app()
        nombres = leer_lista_estudiantes(app.ruta_curso)
        if not 1 <= pos <= len(nombres):
            return
        cont = BoxLayout(orientation="vertical", spacing=dp(8), padding=dp(10))
        cont.add_widget(_caption(f"Código {codigo_de(pos)}"))
        ti = TextInput(text=nombres[pos - 1], multiline=False, size_hint_y=None, height=dp(46))
        cont.add_widget(ti)
        cont.add_widget(Widget())
        fila = BoxLayout(size_hint_y=None, height=dp(46), spacing=dp(8))
        btn_borrar = Factory.BotonPeligro(text="Eliminar")
        btn_ok = Factory.BotonAccion(text="Guardar")
        fila.add_widget(btn_borrar)
        fila.add_widget(btn_ok)
        cont.add_widget(fila)
        pop = Popup(title="Editar estudiante", content=cont, size_hint=(0.9, 0.42))

        def _guardar(*_):
            nuevo = ti.text.strip()
            if nuevo:
                nombres[pos - 1] = nuevo
                guardar_lista_estudiantes(app.ruta_curso, nombres)
            pop.dismiss()
            self.construir()

        def _eliminar():
            del nombres[pos - 1]
            guardar_lista_estudiantes(app.ruta_curso, nombres)
            self.construir()

        def _pedir_confirmacion(*_):
            pop.dismiss()
            siguientes = len(nombres) - pos
            extra = (f"\n\nLos {siguientes} estudiantes que siguen en la lista cambiarán de "
                     f"código (el {codigo_de(pos + 1)} pasa a ser {codigo_de(pos)}, etc.). "
                     "Si ya calificaste exámenes, sus hojas podrían quedar asignadas a otro alumno."
                     if siguientes else "")
            confirmar("Eliminar estudiante",
                      f"¿Eliminar a {nombres[pos - 1]} ({codigo_de(pos)})?{extra}",
                      "Eliminar", _eliminar)

        btn_ok.bind(on_release=_guardar)
        ti.bind(on_text_validate=_guardar)
        btn_borrar.bind(on_release=_pedir_confirmacion)
        pop.open()

    def _vaciar(self, *_):
        app = App.get_running_app()
        n = len(leer_lista_estudiantes(app.ruta_curso))
        if not n:
            return

        def _hacer():
            guardar_lista_estudiantes(app.ruta_curso, [])
            self.construir()
        confirmar("Vaciar lista", f"¿Borrar los {n} estudiantes de la lista de este curso?",
                  "Vaciar", _hacer)

    def _importar(self, *_):
        PF.elegir_archivos(self._planilla_elegida, tipo="planillas", multiple=False,
                           titulo="Elige la planilla (Excel o CSV)")

    def _planilla_elegida(self, refs):
        app = App.get_running_app()
        nombre = PF.nombre_de(refs[0])
        if Path(nombre).suffix.lower() not in (".xlsx", ".xlsm", ".csv"):
            aviso("Formato no válido", f"'{nombre}' no es un Excel (.xlsx) ni un CSV.\n"
                  "Si la lista está en Google Sheets, descárgala como .xlsx.")
            return
        try:
            nombres = importar_planilla(PF.copiar_a_temporal(refs[0]))
        except Exception as e:
            aviso("No se pudo leer", f"No se pudo leer la planilla:\n{e}")
            return
        if not nombres:
            aviso("Sin nombres", "No encontré nombres en ese archivo.\n"
                  "Debe tener una columna con los nombres.")
            return

        def _hacer():
            guardar_lista_estudiantes(app.ruta_curso, nombres)
            self.construir()
            aviso("Listo", f"Se importaron {len(nombres)} estudiantes.")

        actuales = len(leer_lista_estudiantes(app.ruta_curso))
        if actuales:
            confirmar("Reemplazar lista",
                      f"La planilla tiene {len(nombres)} estudiantes y reemplazará la lista "
                      f"actual ({actuales}). Los códigos se asignan según el orden de la planilla.",
                      "Reemplazar", _hacer, peligro=False)
        else:
            _hacer()

# ─────────────────────────────────────────────────────────────────────────────
# PANTALLA: EXÁMENES (Paso 3)
# ─────────────────────────────────────────────────────────────────────────────
class ExamenesScreen(Screen):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.cont = BoxLayout(orientation="vertical", padding=dp(14), spacing=dp(8))
        self.add_widget(self.cont)

    def on_pre_enter(self, *_):
        self.construir()

    def construir(self):
        app = App.get_running_app()
        self.cont.clear_widgets()
        b = self.cont
        b.add_widget(Encabezado(f"Ex\u00e1menes \u00b7 {app.grado}\u00b0{app.curso}"))
        examenes = listar_examenes(app.ruta_curso)
        lista_box = BoxLayout(orientation="vertical", size_hint_y=None, spacing=dp(2))
        lista_box.bind(minimum_height=lista_box.setter("height"))
        if examenes:
            for ruta, nombre, npreg, formato in examenes:
                fila = BoxLayout(size_hint_y=None, height=dp(60), spacing=dp(8))
                tipo = "media hoja" if formato == LH.MEDIA_CARTA.clave else "hoja completa"
                lbl = Label(
                    text=f"{nombre}\n[size=12sp][color=6E7681]{npreg} preguntas"
                         f" · {tipo}[/color][/size]",
                    markup=True, halign="left", valign="middle")
                lbl.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, i.height)))
                fila.add_widget(lbl)
                btn_dl = Factory.BotonOutline(text="Descargar", size_hint_x=None,
                                              width=dp(120))
                btn_dl.bind(on_release=lambda inst, n=nombre, q=npreg, f=formato:
                            self._descargar(n, q, f))
                fila.add_widget(btn_dl)
                btn_x = Factory.BotonPeligro(text="X", size_hint_x=None, width=dp(46))
                btn_x.bind(on_release=lambda inst, r=ruta: self._borrar(r))
                fila.add_widget(btn_x)
                lista_box.add_widget(fila)
        else:
            lista_box.add_widget(Label(text="(Aún no hay exámenes)",
                                       size_hint_y=None, height=dp(30)))
        scroll = ScrollView()
        scroll.add_widget(lista_box)
        b.add_widget(scroll)
        btn_nuevo = Button(text="+ Crear examen nuevo", size_hint_y=None, height=dp(50))
        btn_nuevo.bind(on_release=lambda *_: self._popup_nuevo())
        b.add_widget(btn_nuevo)
        btn_volver = Button(text="← Volver al menú", size_hint_y=None, height=dp(44))
        btn_volver.bind(on_release=lambda *_: setattr(App.get_running_app().sm, "current", "dashboard"))
        b.add_widget(btn_volver)

    def _borrar(self, ruta):
        try:
            nombre = json.loads(ruta.read_text(encoding="utf-8")).get("examen", ruta.stem)
        except Exception:
            nombre = ruta.stem

        def _hacer():
            try:
                ruta.unlink()
            except Exception:
                pass
            self.construir()
        confirmar("Eliminar examen",
                  f"¿Eliminar el examen '{nombre}' y su clave de respuestas?\n"
                  "Las notas ya calificadas no se borran.", "Eliminar", _hacer)

    def _descargar(self, nombre, npreg, formato):
        try:
            cantidad = int(npreg)
        except (ValueError, TypeError):
            aviso("No se pudo", "Ese examen no tiene un número de preguntas válido.")
            return
        try:
            texto, ref, arch = generar_pdf_a_descargas(nombre, cantidad, formato)
            aviso_archivo("Descargado", f"PDF guardado en:\n{texto}", ref, arch)
        except Exception as e:
            aviso("No se pudo descargar", str(e))

    def _popup_nuevo(self):
        cont = BoxLayout(orientation="vertical", spacing=dp(6), padding=dp(10))
        cont.add_widget(Label(text="Nuevo examen", size_hint_y=None, height=dp(28), font_size="16sp"))
        cont.add_widget(Label(text="Nombre:", size_hint_y=None, height=dp(22)))
        in_nombre = TextInput(hint_text="ej. Ingles", multiline=False, size_hint_y=None, height=dp(42))
        cont.add_widget(in_nombre)
        cont.add_widget(Label(text="Número de preguntas:", size_hint_y=None, height=dp(22)))
        in_num = TextInput(hint_text="ej. 10", multiline=False, input_filter="int",
                           size_hint_y=None, height=dp(42))
        cont.add_widget(in_num)
        cont.add_widget(Label(text="Clave (letras seguidas, ej. ABDC):", size_hint_y=None, height=dp(22)))
        in_clave = TextInput(hint_text="ABCD...", multiline=False, size_hint_y=None, height=dp(42))
        cont.add_widget(in_clave)
        cont.add_widget(Label(text="Tipo de hoja:", size_hint_y=None, height=dp(22)))
        formatos = {f.descripcion: f for f in (LH.MEDIA_CARTA, LH.CARTA)}
        sp_formato = Spinner(text=LH.MEDIA_CARTA.descripcion, values=list(formatos),
                             size_hint_y=None, height=dp(42))
        cont.add_widget(sp_formato)

        def _sugerir_formato(_inst, texto):
            # Si no cabe en media hoja, pasar solo a hoja completa
            try:
                n = int(texto)
            except ValueError:
                return
            if n > LH.MEDIA_CARTA.max_preguntas:
                sp_formato.text = LH.CARTA.descripcion
        in_num.bind(text=_sugerir_formato)
        fila_chk = BoxLayout(size_hint_y=None, height=dp(36), spacing=dp(6))
        chk = CheckBox(size_hint_x=None, width=dp(36), active=True)
        fila_chk.add_widget(chk)
        fila_chk.add_widget(Label(text="Guardar el PDF de la hoja en Descargas"))
        cont.add_widget(fila_chk)
        lbl_msg = Label(text="", size_hint_y=None, height=dp(24), color=(0.78, 0.16, 0.16, 1))
        cont.add_widget(lbl_msg)
        btn_guardar = Factory.BotonAccion(text="Guardar", size_hint_y=None, height=dp(48))
        cont.add_widget(btn_guardar)
        pop = Popup(title="Crear examen", content=cont, size_hint=(0.95, 0.9))

        def _guardar(*_):
            app = App.get_running_app()
            nombre = in_nombre.text.strip()
            if not nombre:
                lbl_msg.text = "Escribe el nombre del examen."
                return
            try:
                cantidad = int(in_num.text)
            except ValueError:
                lbl_msg.text = "El número de preguntas debe ser un número."
                return
            formato = formatos[sp_formato.text]
            if not (1 <= cantidad <= formato.max_preguntas):
                lbl_msg.text = (f"{formato.descripcion}: de 1 a "
                                f"{formato.max_preguntas} preguntas.")
                return
            clave = in_clave.text.strip().upper()
            if len(clave) != cantidad or any(c not in "ABCD" for c in clave):
                lbl_msg.text = f"La clave debe tener {cantidad} letras (solo A, B, C, D)."
                return
            existente = app.ruta_curso / "examenes_claves" / f"respuestas_{nombre.replace(' ', '_')}.json"
            if existente.exists():
                lbl_msg.text = "Ya existe un examen con ese nombre."
                return
            guardar_clave(app.ruta_curso, nombre, cantidad, list(clave), formato.clave)
            mensaje = f"Examen '{nombre}' guardado."
            pop.dismiss()
            self.construir()
            if chk.active:
                try:
                    texto, ref, arch = generar_pdf_a_descargas(nombre, cantidad, formato.clave)
                    aviso_archivo("Listo", f"{mensaje}\nPDF en Descargas:\n{texto}", ref, arch)
                    return
                except Exception as e:
                    mensaje += f"\n(No se pudo guardar el PDF: {e})"
            aviso("Listo", mensaje)

        btn_guardar.bind(on_release=_guardar)
        pop.open()


# ─────────────────────────────────────────────────────────────────────────────
# PANTALLA: CALIFICAR (Paso 4)  ·  digitalizar -> leer -> guardar notas
# ─────────────────────────────────────────────────────────────────────────────
class CalificarScreen(Screen):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.cont = BoxLayout(orientation="vertical", padding=dp(14), spacing=dp(8))
        self.add_widget(self.cont)
        self.examen = ""
        self.errores = []     # hojas que no se pudieron leer en esta sesión (no se guardan)

    def on_pre_enter(self, *_):
        self.errores = []
        self.construir()

    @property
    def resultados(self):
        """Lo que se muestra: errores de esta sesión + historial guardado del examen."""
        guardados = sorted(self._leer_notas(), key=lambda r: r.get("fecha", ""), reverse=True)
        return self.errores + guardados

    def construir(self):
        app = App.get_running_app()
        self.cont.clear_widgets()
        b = self.cont
        b.add_widget(Encabezado(f"Calificar · {app.grado}°{app.curso}"))

        examenes = listar_examenes(app.ruta_curso)
        if not examenes:
            b.add_widget(Label(text="Primero crea un examen en 'Exámenes'."))
            self._volver(b)
            return

        self._mapa_claves = {nombre: ruta for ruta, nombre, _, _ in examenes}
        nombres = list(self._mapa_claves.keys())
        if self.examen not in nombres:
            self.examen, self.errores = nombres[0], []
        fila = BoxLayout(size_hint_y=None, height=dp(44), spacing=dp(6))
        fila.add_widget(Label(text="Examen:", size_hint_x=0.3))
        self.sp_examen = Spinner(text=self.examen, values=nombres)
        self.sp_examen.bind(text=self._cambiar_examen)
        fila.add_widget(self.sp_examen)
        b.add_widget(fila)

        btn = Factory.BotonAccion(text="Cargar foto(s) y calificar", size_hint_y=None,
                                  height=dp(52))
        btn.bind(on_release=self._calificar)
        b.add_widget(btn)

        btn_cam = Button(text="Tomar foto con la cámara", size_hint_y=None, height=dp(48))
        btn_cam.bind(on_press=lambda *_: _vibrar())
        btn_cam.bind(on_release=self._tomar_foto)
        b.add_widget(btn_cam)

        lista = BoxLayout(orientation="vertical", size_hint_y=None, spacing=dp(4))
        lista.bind(minimum_height=lista.setter("height"))
        resultados = self.resultados
        if resultados:
            import registro_notas as RN
            # Códigos usados por hojas distintas en TODO el examen
            repetidos = RN.detectar_conflictos(self._leer_notas())
            hojas = len(resultados) - len(self.errores)
            b.add_widget(_caption(f"{hojas} hoja{'s' if hojas != 1 else ''} calificada"
                                  f"{'s' if hojas != 1 else ''}  ·  la más reciente arriba"))
            for r in resultados:
                lista.add_widget(self._fila_resultado(r, repetidos))
        else:
            lista.add_widget(Label(text="Aún no hay hojas calificadas de este examen.",
                                   size_hint_y=None, height=dp(28), font_size="13sp",
                                   color=(0.42, 0.47, 0.55, 1)))
        scroll = ScrollView()
        scroll.add_widget(lista)
        b.add_widget(scroll)

        self._volver(b)

    def _fila_resultado(self, r, repetidos):
        codigo = str(r.get("codigo", "???"))
        if codigo == "ERROR":
            texto = (f"[color={ROJO}][b]No se pudo leer[/b]  ·  "
                     f"{r.get('error') or r['archivo']}[/color]")
        elif not codigo.isdigit():
            texto = f"[color={AMBAR}][b]{codigo}[/b]  ·  código ilegible, corrígelo[/color]"
        elif codigo in repetidos:
            texto = f"[color={AMBAR}][b]{codigo}[/b]  ·  código repetido, revisa[/color]"
        else:
            n = nota_local(r["aciertos"], r["total"])
            col = VERDE if n >= 6.0 else ROJO
            texto = (f"[b]{codigo}[/b]  ·  {r['aciertos']}/{r['total']}"
                     f"  ·  nota [color={col}][b]{n:.1f}[/b][/color]")
        fila = BoxLayout(size_hint_y=None, height=dp(42), spacing=dp(6))
        lbl = Label(text=texto, markup=True, halign="left", valign="middle", font_size="14sp",
                    shorten=True, shorten_from="right")
        lbl.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, i.height)))
        fila.add_widget(lbl)
        btn = Factory.BotonOutline(text="Ver", size_hint_x=None, width=dp(64))
        btn.bind(on_release=lambda inst, rr=r: self._ver_hoja(rr))
        fila.add_widget(btn)
        return fila

    def _cambiar_examen(self, _sp, nombre):
        if nombre != self.examen:
            self.examen, self.errores = nombre, []
            self.construir()

    def _volver(self, b):
        btn = Button(text="← Volver al menú", size_hint_y=None, height=dp(44))
        btn.bind(on_release=lambda *_: setattr(App.get_running_app().sm, "current", "dashboard"))
        b.add_widget(btn)

    def _calificar(self, *_):
        PF.elegir_archivos(lambda refs: self._iniciar_calificacion(refs, mover=False),
                           tipo="imagenes", multiple=True,
                           titulo="Elige la(s) foto(s) de las hojas")

    def _ver_hoja(self, r):
        """Detalle de una hoja: imagen marcada + corregir código / eliminar."""
        import lector_omr as L
        codigo = str(r.get("codigo", "???"))
        cont = BoxLayout(orientation="vertical", spacing=dp(8), padding=[0, dp(4), 0, 0])
        img = None
        if r.get("imagen"):
            img = L.leer_imagen(self._ruta_notas().parent / r["imagen"])
        if img is not None:
            cont.add_widget(Image(texture=self._np_a_textura(img), allow_stretch=True,
                                  keep_ratio=True))
        else:
            msg = (r.get("error") or "No se pudo leer la hoja.") if codigo == "ERROR" \
                else "(Esta hoja no tiene imagen guardada.)"
            lbl = Label(text=msg, halign="center", valign="middle")
            lbl.bind(size=lambda i, *_: setattr(i, "text_size", (i.width, None)))
            cont.add_widget(lbl)
        fila = BoxLayout(size_hint_y=None, height=dp(46), spacing=dp(8))
        titulo = (f"{codigo}  ·  {r['aciertos']}/{r['total']}" if codigo != "ERROR"
                  else r.get("archivo", "Hoja"))
        pop = Popup(title=titulo, content=cont, size_hint=(0.96, 0.92))
        if codigo != "ERROR":
            btn_cod = Button(text="Corregir código", font_size="14sp")
            btn_cod.bind(on_release=lambda *_: (pop.dismiss(), self._editar_codigo(r)))
            fila.add_widget(btn_cod)
        btn_del = Factory.BotonPeligro(text="Eliminar" if codigo != "ERROR" else "Quitar",
                                       font_size="14sp", size_hint_x=0.6)
        btn_del.bind(on_release=lambda *_: (pop.dismiss(), self._eliminar_hoja(r)))
        fila.add_widget(btn_del)
        btn_x = Factory.BotonOutline(text="Cerrar", font_size="14sp", size_hint_x=0.6)
        btn_x.bind(on_release=lambda *_: pop.dismiss())
        fila.add_widget(btn_x)
        cont.add_widget(fila)
        pop.open()

    def _eliminar_hoja(self, r):
        if r.get("codigo") == "ERROR":
            self.errores = [e for e in self.errores if e is not r]
            self.construir()
            return

        def _hacer():
            datos = [e for e in self._leer_notas() if e.get("archivo") != r.get("archivo")]
            self._ruta_notas().write_text(json.dumps(datos, indent=2, ensure_ascii=False),
                                          encoding="utf-8")
            if r.get("imagen"):
                try:
                    (self._ruta_notas().parent / r["imagen"]).unlink()
                except Exception:
                    pass
            self.construir()
        confirmar("Eliminar hoja",
                  f"¿Eliminar la hoja con código {r.get('codigo')} ({r['aciertos']}/{r['total']})?\n"
                  "Su nota ya no aparecerá en el registro. Úsalo si la foto estaba repetida "
                  "o era de otro examen.", "Eliminar", _hacer)

    def _iniciar_calificacion(self, archivos, mover=False, al_terminar=None):
        """
        Califica las fotos EN SEGUNDO PLANO, con una ventana de progreso.

        La interfaz no se congela (en Android así no aparece "La app no
        responde"): un hilo procesa las fotos una por una y avisa a la
        pantalla después de cada una. Cancelar detiene el trabajo al terminar
        la hoja en curso; las hojas ya calificadas se guardan.

        mover=True mueve la foto en vez de copiarla (fotos de la cámara).
        al_terminar(registros): se llama al final en lugar de mostrar el resumen.
        """
        import lector_omr as L
        app = App.get_running_app()
        ruta_curso = app.ruta_curso
        clave_json = self._mapa_claves[self.sp_examen.text]
        try:
            datos = L.cargar_datos_examen(clave_json)
        except Exception as e:
            aviso("Clave dañada", f"No se pudo leer la clave del examen:\n{e}")
            return

        cancelar = threading.Event()
        pop = PopupProgreso(
            "Calificando hojas" if len(archivos) > 1 else "Calificando hoja",
            f"{datos.nombre.replace('_', ' ')}  ·  {app.grado}°{app.curso}",
            len(archivos), cancelar.set)
        pop.open()

        dir_crudos = ruta_curso / "examenes_crudos"
        dir_crudos.mkdir(parents=True, exist_ok=True)

        def estado_de(r):
            if r["codigo"] == "ERROR":
                return "error"
            return "ok" if str(r["codigo"]).isdigit() else "revisar"

        def trabajo():
            nuevos = []
            for i, a in enumerate(archivos):
                if cancelar.is_set():
                    break
                try:
                    nombre = PF.nombre_de(a)
                except Exception:
                    nombre = "foto"
                Clock.schedule_once(lambda dt, i=i, n=nombre: pop.empezar(i, n))
                # Copiar con nombre único (dos 'IMG_0001.jpg' ya no se pisan)
                # y calificar SOLO esa foto
                destino = dir_crudos / L.nombre_unico(nombre)
                try:
                    PF.copiar(a, destino, mover=mover)
                    r = L.calificar_foto(destino, clave_json, ruta_curso, datos)
                except Exception as e:
                    r = {"archivo": nombre, "codigo": "ERROR", "aciertos": 0,
                         "total": datos.total_preguntas, "error": str(e)}
                nuevos.append(r)
                Clock.schedule_once(lambda dt, k=i + 1, e=estado_de(r): pop.avance(k, e))
            Clock.schedule_once(lambda dt: self._fin_calificacion(
                nuevos, datos, pop, cancelar.is_set(), al_terminar))

        threading.Thread(target=trabajo, daemon=True).start()

    def _fin_calificacion(self, nuevos, datos, pop, cancelado, al_terminar):
        """De vuelta en la interfaz: guarda las notas y muestra el resumen."""
        import registro_notas as RN
        if datos.nombre != self.examen:
            self.errores = []
        self.examen = datos.nombre
        self._guardar_registros([r for r in nuevos if r["codigo"] != "ERROR"])
        # Las hojas ilegibles no son notas: se muestran solo en esta sesión
        self.errores = [r for r in nuevos if r["codigo"] == "ERROR"] + self.errores
        self.construir()

        if al_terminar is not None:
            pop.dismiss()
            al_terminar(nuevos)
            return
        repetidos = RN.detectar_conflictos(self._leer_notas())
        cuenta = {"ok": 0, "revisar": 0, "error": 0}
        for r in nuevos:
            cod = str(r["codigo"])
            if cod == "ERROR":
                cuenta["error"] += 1
            elif not cod.isdigit() or cod in repetidos:
                cuenta["revisar"] += 1
            else:
                cuenta["ok"] += 1
        pop.terminar(cuenta, cancelado, al_cerrar=lambda: None)

    def _ruta_notas(self):
        app = App.get_running_app()
        dir_salida = app.ruta_curso / "resultados"
        dir_salida.mkdir(parents=True, exist_ok=True)
        return dir_salida / f"notas_{self.examen}.json"

    def _leer_notas(self):
        notas_path = self._ruta_notas()
        if not notas_path.exists():
            return []
        try:
            return json.loads(notas_path.read_text(encoding="utf-8"))
        except Exception:
            return []

    def _guardar_registros(self, nuevos):
        """Agrega/actualiza registros en notas_<examen>.json (clave: archivo).
        Las hojas con ERROR no se guardan: no son una nota (se ven en pantalla)."""
        registros = {r.get("archivo", ""): r for r in self._leer_notas()}
        for r in nuevos:
            registros[r["archivo"]] = r
        self._ruta_notas().write_text(
            json.dumps(list(registros.values()), indent=2, ensure_ascii=False),
            encoding="utf-8")

    def _editar_codigo(self, r):
        cont = BoxLayout(orientation="vertical", spacing=dp(8), padding=dp(10))
        cont.add_widget(Label(text=f"Hoja: {r['archivo']}", size_hint_y=None, height=dp(24)))
        cont.add_widget(Label(text="Código correcto (3 dígitos):",
                              size_hint_y=None, height=dp(24)))
        actual = "".join(c for c in r.get("codigo", "") if c.isdigit())
        ti = TextInput(text=actual, multiline=False, input_filter="int",
                       size_hint_y=None, height=dp(46))
        cont.add_widget(ti)
        lbl = Label(text="", size_hint_y=None, height=dp(22), color=(0.78, 0.16, 0.16, 1))
        cont.add_widget(lbl)
        btn = Factory.BotonAccion(text="Guardar c\u00f3digo", size_hint_y=None, height=dp(46))
        cont.add_widget(btn)
        pop = Popup(title="Corregir código", content=cont, size_hint=(0.85, 0.55))

        def _ok(*_):
            nuevo = ti.text.strip()
            if not nuevo.isdigit():
                lbl.text = "Escribe solo números."
                return
            nuevo = nuevo.zfill(3)
            r["codigo"] = nuevo
            self._actualizar_notas_codigo(r["archivo"], nuevo)
            pop.dismiss()
            self.construir()

        btn.bind(on_release=_ok)
        pop.open()

    def _actualizar_notas_codigo(self, archivo, nuevo):
        data = self._leer_notas()
        for e in data:
            if e.get("archivo") == archivo:
                e["codigo"] = nuevo
                e["codigo_corregido"] = True
        self._ruta_notas().write_text(json.dumps(data, indent=2, ensure_ascii=False),
                                      encoding="utf-8")

    # ── Cámara (solo Android / iOS) ───────────────────────────────────────────
    def _ruta_captura(self):
        """Carpeta donde la cámara guardará la foto. En Android usa la carpeta
        externa de la app (sí es accesible por la app de cámara); en otro caso,
        el directorio de datos del usuario."""
        if platform == "android":
            try:
                from jnius import autoclass  # type: ignore
                PythonActivity = autoclass("org.kivy.android.PythonActivity")
                ctx = PythonActivity.mActivity
                ext = ctx.getExternalFilesDir(None)
                if ext is not None:
                    return Path(ext.getAbsolutePath()) / "captura_omr.jpg"
            except Exception:
                pass
        return Path(App.get_running_app().user_data_dir) / "captura_omr.jpg"

    def _tomar_foto(self, *_):
        if platform == "android":
            self._tomar_foto_android()
            return
        if platform == "ios":
            try:
                from plyer import camera  # type: ignore
            except Exception:
                aviso("Cámara", "No se pudo acceder a la cámara en este dispositivo.")
                return
            destino = self._ruta_captura()
            try:
                destino.parent.mkdir(parents=True, exist_ok=True)
            except Exception:
                pass
            try:
                camera.take_picture(filename=str(destino), on_complete=self._foto_lista)
            except Exception as e:
                aviso("Cámara", f"No se pudo abrir la cámara: {e}")
            return
        aviso("Cámara", "Función solo disponible en dispositivos android o ios.")

    # Código de petición para identificar el resultado de la cámara
    _COD_CAMARA = 0xC0FE

    def _tomar_foto_android(self):
        """Abre la cámara con el Intent nativo de Android y espera el resultado.
        Es más confiable que plyer: el sistema nos avisa cuando la foto está
        lista mediante on_activity_result."""
        try:
            from jnius import autoclass, cast  # type: ignore
            from android import activity  # type: ignore
        except Exception as e:
            aviso("Cámara", f"No se pudo acceder a la cámara: {e}")
            return
        # Permite pasar la ruta file:// a la app de cámara sin que Android cierre la app.
        try:
            StrictMode = autoclass("android.os.StrictMode")
            StrictMode.disableDeathOnFileUriExposure()
        except Exception:
            pass

        destino = self._ruta_captura()
        try:
            destino.parent.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
        self._destino_foto = destino

        try:
            PythonActivity = autoclass("org.kivy.android.PythonActivity")
            Intent = autoclass("android.content.Intent")
            MediaStore = autoclass("android.provider.MediaStore")
            Uri = autoclass("android.net.Uri")
            File = autoclass("java.io.File")

            uri = Uri.fromFile(File(str(destino)))
            intent = Intent(MediaStore.ACTION_IMAGE_CAPTURE)
            intent.putExtra(MediaStore.EXTRA_OUTPUT, cast("android.os.Parcelable", uri))

            activity.bind(on_activity_result=self._on_cam_result)
            PythonActivity.mActivity.startActivityForResult(intent, self._COD_CAMARA)
        except Exception as e:
            try:
                activity.unbind(on_activity_result=self._on_cam_result)
            except Exception:
                pass
            aviso("Cámara", f"No se pudo abrir la cámara: {e}")

    def _on_cam_result(self, request, result, intent):
        if request != self._COD_CAMARA:
            return
        try:
            from android import activity  # type: ignore
            activity.unbind(on_activity_result=self._on_cam_result)
        except Exception:
            pass
        # RESULT_OK vale -1 en Android; si el usuario cancela, no hacemos nada.
        if result != -1:
            return
        destino = getattr(self, "_destino_foto", None)
        ruta = str(destino) if destino else ""
        # Procesamos en el hilo principal de Kivy (UI segura) tras un instante,
        # dándole tiempo a la cámara de terminar de escribir el archivo.
        Clock.schedule_once(lambda dt: self._foto_lista(ruta), 0.3)

    def _foto_lista(self, ruta_foto):
        """La cámara terminó: se califica la foto con la misma ventana de
        progreso que las fotos cargadas, y luego se muestra la hoja marcada."""
        p = Path(str(ruta_foto)) if ruta_foto else None
        if not p or not p.exists():
            aviso("Cámara", "No se recibió la foto.")
            return
        # Se mueve a la carpeta del curso (no queda en la galería)
        self._iniciar_calificacion([p], mover=True, al_terminar=self._mostrar_foto_calificada)

    def _mostrar_foto_calificada(self, registros):
        r = registros[0] if registros else None
        if r is None or r["codigo"] == "ERROR":
            aviso("No se pudo calificar", (r or {}).get("error") or "No se pudo leer la hoja.")
            return
        self._ver_hoja(r)

    @staticmethod
    def _np_a_textura(img_bgr):
        import cv2
        rgb = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2RGB)
        rgb = cv2.flip(rgb, 0)  # el origen de la textura de Kivy está abajo
        h, w = rgb.shape[:2]
        tex = Texture.create(size=(w, h), colorfmt="rgb")
        tex.blit_buffer(rgb.tobytes(), colorfmt="rgb", bufferfmt="ubyte")
        return tex


# ─────────────────────────────────────────────────────────────────────────────
# PANTALLA: REGISTRO DE NOTAS (Paso 5)  ·  Excel cruzando código -> estudiante
# ─────────────────────────────────────────────────────────────────────────────
class RegistroScreen(Screen):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.cont = BoxLayout(orientation="vertical", padding=dp(14), spacing=dp(8))
        self.add_widget(self.cont)
        self.ultimo_excel = None
        self.mensaje = ""

    def on_pre_enter(self, *_):
        self.ultimo_excel = None
        self.mensaje = ""
        self.construir()

    def construir(self):
        app = App.get_running_app()
        self.cont.clear_widgets()
        b = self.cont
        b.add_widget(Encabezado(f"Registro de notas \u00b7 {app.grado}\u00b0{app.curso}"))

        import registro_notas as RN
        examenes = RN.examenes_disponibles(app.ruta_curso)
        if not examenes:
            b.add_widget(Label(text="Aún no hay exámenes calificados.\n"
                                    "Primero califica hojas en 'Calificar'."))
            self._volver(b)
            return

        fila = BoxLayout(size_hint_y=None, height=dp(44), spacing=dp(6))
        fila.add_widget(Label(text="Examen:", size_hint_x=0.3))
        self.sp_examen = Spinner(text=examenes[0], values=examenes)
        fila.add_widget(self.sp_examen)
        b.add_widget(fila)

        btn = Factory.BotonAccion(text="Generar registro Excel", size_hint_y=None,
                                  height=dp(52))
        btn.bind(on_release=self._generar)
        b.add_widget(btn)

        self.lbl_msg = Label(text=self.mensaje, size_hint_y=None, height=dp(60))
        b.add_widget(self.lbl_msg)

        if self.ultimo_excel:
            ref, nombre = self.ultimo_excel
            fila_b = BoxLayout(size_hint_y=None, height=dp(46), spacing=dp(8))
            btn_abrir = Button(text="Abrir Excel")
            btn_abrir.bind(on_release=lambda *_: self._abrir(ref, nombre))
            fila_b.add_widget(btn_abrir)
            btn_comp = Factory.BotonOutline(text="Compartir" if PF.ES_ANDROID else "Ver carpeta")
            btn_comp.bind(on_release=lambda *_: self._compartir(ref, nombre))
            fila_b.add_widget(btn_comp)
            b.add_widget(fila_b)

        b.add_widget(Label(text="", size_hint_y=1))
        self._volver(b)

    def _volver(self, b):
        btn = Button(text="← Volver al menú", size_hint_y=None, height=dp(44))
        btn.bind(on_release=lambda *_: setattr(App.get_running_app().sm, "current", "dashboard"))
        b.add_widget(btn)

    def _generar(self, *_):
        app = App.get_running_app()
        import registro_notas as RN
        examen = self.sp_examen.text
        try:
            ruta = RN.generar_registro(app.ruta_curso, examen)
        except Exception as e:
            aviso("Error", f"No se pudo generar el registro:\n{e}")
            return
        if ruta is None:
            aviso("Falta la planilla",
                  "No hay lista de estudiantes para este curso.\n"
                  "Créala en 'Estudiantes' y vuelve a intentar.")
            return
        # Copia en Descargas para encontrarlo, abrirlo o compartirlo fuera de la app
        try:
            texto, ref = PF.guardar_en_descargas(ruta, ruta.name)
        except Exception:
            texto, ref = str(ruta), str(ruta)
        self.ultimo_excel = (ref, ruta.name)
        self.mensaje = f"Registro generado:\n{texto}"
        self.construir()

    def _abrir(self, ref, nombre):
        err = PF.abrir(ref, nombre)
        if err:
            aviso("No se pudo abrir",
                  f"Ábrelo desde Descargas ({nombre}).\n"
                  "Necesitas una app de hojas de cálculo (Excel, Google Sheets o WPS).\n\n"
                  f"({err})")

    def _compartir(self, ref, nombre):
        err = PF.compartir(ref, nombre)
        if err:
            aviso("No se pudo compartir", err)


# ─────────────────────────────────────────────────────────────────────────────
# APP
# ─────────────────────────────────────────────────────────────────────────────
class CalificadorApp(App):
    profesor = ""
    grado = ""
    curso = ""
    ruta_curso = None

    def build(self):
        self.title = "CalificadorApp"
        # En Android la carpeta de instalación es de solo lectura: movemos el
        # "directorio de trabajo" a la carpeta privada de la app (sí escribible),
        # para que 'profesores/' y todos los datos se guarden sin problemas.
        # En PC, los datos van siempre junto a la app (antes dependían de la
        # carpeta desde donde se abría, y parecía que se "perdían").
        try:
            os.chdir(self.user_data_dir if platform == "android" else _DIR)
        except Exception:
            pass
        _registrar_fuentes()
        Window.clearcolor = FONDO
        if platform != "android":
            Window.size = (400, 720)   # vista tipo celular para probar en PC
        GS.inicializar_sistema()
        self.sm = ScreenManager(transition=SlideTransition(duration=0.18))
        self.sm.add_widget(SetupScreen(name="setup"))
        self.sm.add_widget(HomeScreen(name="home"))
        self.sm.add_widget(DashboardScreen(name="dashboard"))
        self.sm.add_widget(EstudiantesScreen(name="estudiantes"))
        self.sm.add_widget(ExamenesScreen(name="examenes"))
        self.sm.add_widget(CalificarScreen(name="calificar"))
        self.sm.add_widget(RegistroScreen(name="registro"))
        self.sm.current = "home" if _subdirs(BASE) else "setup"
        Window.bind(on_keyboard=self._tecla)
        return self.sm

    # Pantalla a la que lleva "Atrás" desde cada una
    _ATRAS = {"estudiantes": "dashboard", "examenes": "dashboard",
              "calificar": "dashboard", "registro": "dashboard", "dashboard": "home"}

    def _tecla(self, _window, tecla, *_):
        """Botón "Atrás" de Android (y Esc en PC): vuelve a la pantalla anterior
        en lugar de cerrar la app. Solo en el inicio cierra la app."""
        if tecla != 27:
            return False
        if any(isinstance(w, ModalView) for w in Window.children):
            return True   # hay una ventana abierta (p. ej. calificando): no salir
        destino = self._ATRAS.get(self.sm.current)
        if destino:
            self.sm.current = destino
            return True
        return False

    def on_start(self):
        _solicitar_permisos_android()

    def ir_a_home(self):
        self.sm.current = "home"


if __name__ == "__main__":
    CalificadorApp().run()